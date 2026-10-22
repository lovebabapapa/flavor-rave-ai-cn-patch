# -*- coding: utf-8 -*-
"""导出全部汉化内容为 Excel 核对表（英文原文 ↔ 中文译文）。"""
import json, os, glob, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
PATCH = os.path.join(HERE, "汉化输出", "FR_CN_补丁")
SRC = os.path.join(HERE, "source_files")
OUT = os.path.join(HERE, "汉化核对表.xlsx")

HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
ALT_FILL = PatternFill("solid", fgColor="F2F6FC")


def src_path(rel):
    if rel.startswith("assets/"):
        return os.path.join(SRC, "assets", "preload", rel[len("assets/"):].replace("/", os.sep))
    return os.path.join(SRC, rel.replace("/", os.sep))


def load(p):
    try:
        return json.load(open(p, encoding="utf-8-sig"))
    except Exception:
        return None


def walk_files(root):
    out = []
    for r, _, fs in os.walk(root):
        for f in fs:
            if f.endswith(".json"):
                out.append(os.path.relpath(os.path.join(r, f), root).replace(os.sep, "/"))
    return sorted(out)


def dialogue_rows():
    rows = []
    for rel in walk_files(PATCH):
        if not rel.endswith(".json"):
            continue
        zh = load(os.path.join(PATCH, rel.replace("/", os.sep)))
        en = load(src_path(rel))
        if not isinstance(zh, dict) or not isinstance(en, dict):
            continue
        ze = zh.get("dialogue") or []
        ee = en.get("dialogue") or []
        # 分类与章节名
        if "/synsun/" in rel:
            kind = "电视百科"
            chapter = os.path.basename(rel)[:-5]
        elif "/images/results/character/" in rel:
            kind = "结算台词"
            chapter = rel.split("/")[-2]
        elif rel.endswith(("dreamcast.json", "saturn.json")):
            kind = "剧情"
            parts = rel.split("/")
            if rel.startswith("assets/data/"):
                chapter = parts[2]
            elif "/data/" in rel:
                chapter = parts[parts.index("data") + 1] if "data" in parts else os.path.basename(rel)[:-5]
            else:
                chapter = os.path.basename(rel)[:-5]
        else:
            continue
        for i, (a, b) in enumerate(zip(ze, ee)):
            if not isinstance(a, dict) or not isinstance(b, dict):
                continue
            zt = (a.get("text") or "").strip()
            et = (b.get("text") or "").strip()
            if not et and not zt:
                continue
            c = a.get("chara") or b.get("chara")
            who = c[0] if isinstance(c, list) and c else (c if isinstance(c, str) else "")
            rows.append([kind, chapter, rel, i, who, et, zt])
    return rows


def flavor_rows():
    rows = []
    for rel in walk_files(PATCH):
        base = os.path.basename(rel)
        if base not in ("flavorpediaData.json", "flavorpediaSideData.json"):
            continue
        zh = load(os.path.join(PATCH, rel.replace("/", os.sep)))
        en = load(src_path(rel))
        if not isinstance(zh, dict) or not isinstance(en, dict):
            continue
        for key in ("character", "place", "things"):
            for i, (a, b) in enumerate(zip(zh.get(key) or [], en.get(key) or [])):
                if not isinstance(a, dict) or not isinstance(b, dict):
                    continue
                ident = a.get("chara") or a.get("name") or b.get("chara") or b.get("name") or f"{key}[{i}]"
                for fld in ("flavorText", "subtitles", "design", "desc", "place"):
                    zt = (a.get(fld) or "").strip()
                    et = (b.get(fld) or "").strip()
                    if not et and not zt:
                        continue
                    rows.append([rel, ident, key, fld, et, zt])
    return rows


def achievements_rows():
    rows = []
    for rel in walk_files(PATCH):
        if not rel.endswith("achievements.json"):
            continue
        zh = load(os.path.join(PATCH, rel.replace("/", os.sep)))
        en = load(src_path(rel))
        if not isinstance(zh, dict) or not isinstance(en, dict):
            continue
        for a, b in zip(zh.get("achievements") or [], en.get("achievements") or []):
            if not isinstance(a, dict) or not isinstance(b, dict):
                continue
            for fld in ("fancyName", "unlockCondition", "showmanDialogue"):
                rows.append([rel, b.get("name", ""), fld, b.get(fld, ""), a.get(fld, "")])
    return rows


def gallery_credits_rows():
    rows = []
    for rel in walk_files(PATCH):
        base = os.path.basename(rel)
        zh = load(os.path.join(PATCH, rel.replace("/", os.sep)))
        en = load(src_path(rel))
        if not isinstance(zh, dict) or not isinstance(en, dict):
            continue
        if base in ("galleryCGData.json", "galleryData.json"):
            for a, b in zip(zh.get("artwork") or [], en.get("artwork") or []):
                if isinstance(a, dict) and isinstance(b, dict):
                    rows.append([base, b.get("image", ""), "caption", b.get("caption", ""), a.get("caption", "")])
        elif base == "creditsData.json":
            for a, b in zip(zh.get("peeps") or [], en.get("peeps") or []):
                if isinstance(a, dict) and isinstance(b, dict):
                    rows.append([base, b.get("realName", ""), "description", b.get("description", ""), a.get("description", "")])
    return rows


def weeks_rows():
    rows = []
    for rel in walk_files(PATCH):
        if "/weeks/" not in rel:
            continue
        zh = load(os.path.join(PATCH, rel.replace("/", os.sep)))
        en = load(src_path(rel))
        if not isinstance(zh, dict) or not isinstance(en, dict):
            continue
        for fld in ("weekName", "storyName", "location"):
            rows.append([rel, fld, en.get(fld, ""), zh.get(fld, "")])
    return rows


def song_title_rows():
    rows = []
    for rel in walk_files(PATCH):
        if not rel.endswith("meta.json"):
            continue
        zh = load(os.path.join(PATCH, rel.replace("/", os.sep)))
        en = load(src_path(rel))
        if not isinstance(zh, dict) or not isinstance(en, dict):
            continue
        zs = zh.get("song") or {}
        es = en.get("song") or {}
        rows.append([rel, es.get("name", ""), zs.get("name", ""), es.get("artist", "")])
    return rows


def lang_rows(cat):
    """语言表：英文原文（若能找到）+ 中文"""
    en_map = {}
    mk = os.path.join(HERE, "_menu_keys.json")
    if os.path.exists(mk):
        d = json.load(open(mk, encoding="utf-8"))
        en_map.update(d.get(cat.lower(), {}))
    ts = os.path.join(HERE, "translate_sheet.json")
    if os.path.exists(ts):
        d = json.load(open(ts, encoding="utf-8"))
        for it in d.get("lyrics", []) + d.get("menu_subtitles", []):
            if it.get("key"):
                en_map[it["key"]] = it.get("en", "")
    p = os.path.join(PATCH, "assets", "data", f"language{cat}.json")
    if not os.path.exists(p):
        return []
    d = json.load(open(p, encoding="utf-8"))
    rows = []
    for it in d.get("map", []):
        k = it.get("key", "")
        v = it.get("value", "")
        if not (v or "").strip():
            continue
        rows.append([k, en_map.get(k, ""), v])
    return rows


def add_sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    headers = list(headers) + ["校对备注"]
    widths = list(widths) + [30]
    rows = [list(r) + [""] for r in rows]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    for i, r in enumerate(rows, 2):
        ws.append(r)
        if i % 2 == 0:
            for c in range(1, len(headers) + 1):
                ws.cell(row=i, column=c).fill = ALT_FILL
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    d = dialogue_rows()
    add_sheet(wb, "剧情与对话", ["分类", "章节/条目", "文件", "序号", "角色", "英文原文", "中文译文"], d, [10, 24, 52, 6, 12, 76, 76])
    for row in wb["剧情与对话"].iter_rows(min_row=2, min_col=6, max_col=7):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    f = flavor_rows()
    add_sheet(wb, "图鉴百科", ["文件", "条目", "分类", "字段", "英文原文", "中文译文"], f, [46, 16, 10, 12, 70, 70])

    a = achievements_rows()
    add_sheet(wb, "成就", ["文件", "ID", "字段", "英文原文", "中文译文"], a, [46, 20, 16, 60, 60])

    g = gallery_credits_rows()
    add_sheet(wb, "画廊与职员表", ["文件", "对象", "字段", "英文原文", "中文译文"], g, [22, 18, 14, 60, 60])

    w = weeks_rows()
    add_sheet(wb, "周目与地点", ["文件", "字段", "英文原文", "中文译文"], w, [52, 12, 46, 46])

    s = song_title_rows()
    add_sheet(wb, "曲名", ["文件", "英文曲名", "中文显示名", "曲作者"], s, [52, 28, 34, 34])

    o = lang_rows("Option")
    add_sheet(wb, "选项菜单", ["key", "英文原文", "中文译文"], o, [34, 56, 56])

    gp = lang_rows("Gameplay")
    add_sheet(wb, "界面文本", ["key", "英文原文", "中文译文"], gp, [34, 46, 56])

    fl = lang_rows("Flavor")
    add_sheet(wb, "字幕与歌词", ["key", "英文原文", "中文译文"], fl, [34, 66, 66])

    info = wb.create_sheet("说明", 0)
    info.column_dimensions["A"].width = 100
    lines = [
        "Flavor Rave 汉化核对表",
        "",
        "用法：各 sheet 均为「英文原文 ↔ 中文译文」对照，最后一列「校对备注」留空供你填写。",
        "改完把这张表发回即可，我按 key/文件+序号写回补丁。",
        "",
        "sheet 说明：",
        "  剧情与对话：剧情过场（dreamcast/saturn）、电视百科（synsun）、结算台词",
        "  图鉴百科：Flavorpedia 角色/地点/物品的正文、字幕、设计者",
        "  成就：名称、解锁条件、Showman 台词",
        "  画廊与职员表：CG 说明、Credits 职位描述",
        "  周目与地点：周目名、章节名、地点",
        "  曲名：「英文名（中文名）」显示格式 + 曲作者",
        "  选项菜单 / 界面文本 / 字幕与歌词：UI 语言表（含海盗船歌 45 句歌词、主菜单字幕）",
        "",
        "注意：",
        "  · 字幕类文本含 :dur:数字 与 [NL]，改动时请保留，否则节奏会乱。",
        "  · 角色名、曲名（英文部分）、地名按既定规则处理，校对时留意不要混入中文音译。",
    ]
    for i, t in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=14)
            c.fill = HEAD_FILL
            c.font = Font(bold=True, size=14, color="FFFFFF")

    wb.save(OUT)
    total = sum(ws.max_row - 1 for ws in wb.worksheets)
    print(f"已生成 {OUT}")
    for ws in wb.worksheets:
        print(f"  {ws.title}: {ws.max_row-1} 行")
    print(f"合计 {total} 行")


if __name__ == "__main__":
    main()
